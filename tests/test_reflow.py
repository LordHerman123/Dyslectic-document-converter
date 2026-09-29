"""The reading view (text flowing to fit the window) and where each document was left."""
from dyslexia_converter import pipeline
from dyslexia_converter.settings import FormatSettings
from dyslexia_converter.ui.app import ConverterApp


def _app_with(paper) -> ConverterApp:
    app = ConverterApp(page=None)
    app.session = pipeline.load(paper, FormatSettings())
    app.source_path = str(paper)
    app.doc_key = "doc-1"
    return app


def test_read_aloud_words_point_at_their_place_in_the_text(paper):
    """Every word handed to the speech engine knows where it is in its paragraph, so the word being said can be
    marked in the text."""
    app = _app_with(paper)
    r = app.reflow
    r.result = app.session.compose(app.settings)
    r.items = r.result.items
    r._sentences()
    assert r._units and len(r._units) == len(r._unit_spans)
    for unit, spans in zip(r._units, r._unit_spans):
        text = r.plain(r.items[unit.page])
        for word, (a, b) in zip(unit.words, spans):
            assert text[a:b] == word.text
    # a sentence the engine gets is the words joined with single spaces, and word starts point into it
    first = r._units[0]
    assert all(first.text[w.start:w.start + len(w.text)] == w.text for w in first.words)


def test_reading_position_is_kept_per_document(paper):
    app = _app_with(paper)
    assert app.reading_position("reflow", 0) == 0
    app.save_reading_position("reflow", 12)
    app.save_reading_position("focus", 3)
    assert app.reading_position("reflow") == 12 and app.reading_position("focus") == 3
    app.doc_key = "doc-2"
    assert app.reading_position("reflow", 0) == 0  # another document starts at the beginning
    # the list stays bounded: the oldest documents are forgotten
    for n in range(app.POSITIONS_KEPT + 5):
        app.doc_key = f"many-{n}"
        app.save_reading_position("focus", n)
    assert len(app.ui["positions"]) == app.POSITIONS_KEPT
    app.doc_key = "doc-1"
    assert app.reading_position("focus", 0) == 0


def test_study_sheet_groups_highlights_under_their_headings():
    """The study sheet lists the highlights under the heading they are in, with pages and notes, and counts them."""
    import io

    from docx import Document

    from dyslexia_converter import highlights as hl

    labels = {"summary": "{highlights} highlights, {notes} with a note", "note": "Note:", "page": "(page {page})",
              "start": "Before the first heading", **{c: c.capitalize() for c in hl.COLOURS}}
    sections = [("", [(1, "yellow", "An opening line", "")]),
                ("1 Introduction", [(2, "green", "Data changes research", "Main claim"),
                                    (3, "yellow", "as boyd argues", "")]),
                ("2 Method", [(5, "pink", "We asked 40 people", "Small sample?")])]
    data = hl.study_sheet_docx("Study sheet: paper", sections, labels)
    d = Document(io.BytesIO(data))
    text = [p.text for p in d.paragraphs]
    assert text[0] == "Study sheet: paper"
    assert "4 highlights, 2 with a note" in text
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading 2")]
    assert heads == ["Before the first heading", "1 Introduction", "2 Method"]
    joined = "\n".join(text)
    assert "“Data changes research”  (page 2)" in joined and "Note: Main claim" in joined
    assert joined.index("Data changes research") < joined.index("We asked 40 people")
    assert "Yellow: 2" in joined


def test_formula_pictures_are_read_as_their_text(paper):
    """A small formula picture inside a line is shown (and read) as its text in the reading view."""
    from dyslexia_converter.model import ImageData
    from dyslexia_converter.render.compose import RItem, Run

    app = _app_with(paper)
    r = app.reflow
    r.result = app.session.compose(app.settings)
    r.result.inline_images = {"": ImageData(b"", "png", 10, 10, kind="equation", alt="x^2")}
    item = RItem("paragraph", [Run("area  grows")])
    assert r.plain(item) == "area  x^2  grows"


def test_colour_help_marks_every_other_syllable_or_sentence(paper):
    """Colour help: the second colour goes on every other syllable of longer words, or on every other sentence;
    the text itself is unchanged (the spans put together are the paragraph)."""
    from dyslexia_converter.render.compose import RItem, Run

    app = _app_with(paper)
    r = app.reflow
    r.result = app.session.compose(app.settings)
    item = RItem("paragraph", [Run("Reading "), Run("wonderful", bold=True), Run(" stories. It helps. Dr. Smith agrees.")])
    r.items = [item]
    text = r.plain(item)

    app.ui["reflow_colours"] = "syllables"
    r._marks = {}
    marks = r._colour_marks(0, item)
    assert text[marks[0][0]:marks[0][1]] == "ing"  # Read-ing: the second syllable
    assert all(text[a:b].isalpha() for a, b in marks)
    spans = r._spans(0, item)
    assert "".join(s.text for s in spans) == text
    assert any(s.style.color for s in spans) and any(not s.style.color for s in spans)
    assert "".join(s.text for s in spans if s.style.weight) == "wonderful"  # bold stays on its run

    app.ui["reflow_colours"] = "sentences"
    r._marks = {}
    marks = r._colour_marks(0, item)
    assert [text[a:b] for a, b in marks] == ["It helps."]  # "Dr." does not end a sentence

    app.ui["reflow_colours"] = "off"
    r._marks = {}
    assert not any(s.style.color for s in r._spans(0, item))


def test_reading_along_marks_sentence_and_word(paper):
    from dyslexia_converter.render.compose import RItem, Run

    app = _app_with(paper)
    r = app.reflow
    r.result = app.session.compose(app.settings)
    item = RItem("paragraph", [Run("One two three. Four five.")])
    r.items = [item]
    spans = r._spans(0, item, lit=(4, 7), sentence=(0, 14))
    by_text = {s.text: s.style.bgcolor for s in spans}
    assert by_text["two"] and by_text["One "] and by_text["two"] != by_text["One "]
    assert by_text[" Four five."] is None


def test_reading_speed_is_kept_per_document(paper):
    app = _app_with(paper)
    app.ui["read_speed"] = 1.25
    assert app.reflow.speed == 1.25  # the app's speed until one is chosen for the document
    app.save_reading_position("speed", 0.75)
    assert app.reflow.speed == 0.75
    app.doc_key = "doc-2"
    assert app.reflow.speed == 1.25
