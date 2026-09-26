import re

import pymupdf
import pytest
from docx import Document as Docx

from conftest import sha256
from dyslexia_converter import pipeline
from dyslexia_converter.model import BlockKind
from dyslexia_converter.settings import PRESETS, FormatSettings


def norm(t: str) -> str:
    return re.sub(r"\s+", " ", t.replace("­", "")).strip()


def pdf_text(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as d:
        return norm(" ".join(p.get_text() for p in d))


@pytest.mark.parametrize("preset", list(PRESETS))
def test_wording_is_preserved(paper, preset):
    session = pipeline.load(paper)
    # no optional transformations and no page furniture interrupting the text flow
    s = PRESETS[preset].copy(move_footnotes=False, page_numbers=False, include_contents=False, about_note=False)
    out = pdf_text(session.export("pdf", s))
    for b in session.document.blocks:
        if b.kind in (BlockKind.PARAGRAPH, BlockKind.HEADING, BlockKind.REFERENCE, BlockKind.CAPTION):
            # every word sequence of the original appears in the output
            assert norm(b.text) in out, b.text[:60]


def test_pdf_is_a4_selectable_with_embedded_fonts_and_outline(paper):
    data = pipeline.load(paper).export("pdf", FormatSettings())
    with pymupdf.open(stream=data, filetype="pdf") as d:
        assert round(d[0].rect.width) == 595 and round(d[0].rect.height) == 842
        embedded = {f[3].split("+")[-1] for p in d for f in p.get_fonts() if f[1] == "ttf"}
        used = {span["font"] for p in d for b in p.get_text("dict")["blocks"] for l in b.get("lines", [])
                for span in l["spans"]}
        assert used and used <= embedded  # all text uses embedded (subset) fonts
        assert any("DejaVu" in f for f in used)
        titles = [t[1] for t in d.get_toc()]
        assert "1 Introduction" in titles and "2.1 Participants" in titles
        assert len(d[1].get_images()) + sum(len(p.get_images()) for p in d) > 0


def test_original_is_never_modified(paper):
    before = sha256(paper)
    session = pipeline.load(paper)
    for fmt in ("pdf", "printable_pdf", "docx", "txt", "md"):
        assert session.export(fmt, FormatSettings())
    assert sha256(paper) == before


def test_citations_moved_to_numbers_and_footnotes_to_notes(paper):
    s = FormatSettings(move_citations=True, move_footnotes=True)
    out = pdf_text(pipeline.load(paper).export("pdf", s))
    assert "change reading speed [4][2]." in out
    assert "(Smith, 2020" not in out
    assert "reader [Note 1]" in out
    assert "[Note 1] Individual differences are discussed in Section 4." in out
    s2 = FormatSettings(move_citations=False)
    assert "(Smith, 2020; Jones & Brown, 2021)" in pdf_text(pipeline.load(paper).export("pdf", s2))


def test_bold_word_start_does_not_change_text(paper):
    plain = pdf_text(pipeline.load(paper).export("pdf", FormatSettings()))
    bionic = pdf_text(pipeline.load(paper).export("pdf", FormatSettings(bold_word_start=True)))
    assert plain == bionic


def test_typography_settings_change_layout(paper):
    session = pipeline.load(paper)
    small = session.export("pdf", FormatSettings(font_size=10, line_spacing=1.2, include_contents=False))
    big = session.export("pdf", FormatSettings(font_size=18, line_spacing=2.2, letter_spacing=1,
                                               word_spacing=4, reading_width=10, include_contents=False))
    count = lambda b: pymupdf.open(stream=b, filetype="pdf").page_count  # noqa: E731
    assert count(big) > count(small)


@pytest.mark.parametrize("font", ["Atkinson Hyperlegible", "OpenDyslexic", "Arial", "Verdana", "Tahoma"])
def test_all_fonts_render(paper, font):
    out = pdf_text(pipeline.load(paper).export("pdf", FormatSettings(font=font, alignment="justify")))
    assert "Reading difficulties affect a substantial proportion" in out


def test_docx_txt_md(paper):
    session = pipeline.load(paper)
    d = Docx(__import__("io").BytesIO(session.export("docx", FormatSettings())))
    assert any(p.style.name.startswith("Heading") and p.text == "1 Introduction" for p in d.paragraphs)
    assert d.tables and d.tables[0].cell(0, 0).text == "Condition"
    txt = session.export("txt", FormatSettings()).decode()
    assert "Reading difficulties affect" in txt
    md = session.export("md", FormatSettings()).decode()
    assert "## 1 Introduction" in md and "| Condition | Time (s) | Accuracy |" in md


def test_unicode_is_preserved(tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    from dyslexia_converter.fonts import font_path

    src = tmp_path / "u.pdf"
    pdfmetrics.registerFont(TTFont("DV", font_path("DejaVu Sans")))
    c = canvas.Canvas(str(src), pagesize=A4)
    c.setFont("DV", 11)
    c.drawString(72, 700, "Greek αβγ and math ≤ ∞, Dutch ëï, ĳsland.")
    c.save()
    out = pdf_text(pipeline.load(src).export("pdf", FormatSettings(font="Atkinson Hyperlegible")))
    assert "Greek αβγ and math ≤ ∞, Dutch ëï, ĳsland." in out


def test_padded_paragraph_never_overflows_a_page():
    """Regression: a one-line boxed quote at the foot of a page made ReportLab fail."""
    from dyslexia_converter.render.compose import Run
    from dyslexia_converter.render.pdf_writer import _styles, RichParagraph

    st = _styles(FormatSettings(), printable=False)["quote"]
    p = RichParagraph([Run("erience.")], st)
    _, h = p.wrap(400, 800)
    assert p.split(400, h - 1) == []  # does not fit: move it, never claim it fits
    assert p.split(400, h + 1) == [p]


def test_added_labels_follow_the_document_language(paper):
    """Words the converter adds (Contents, Notes, [Note n]) use the document's language; the text is unchanged."""
    s = FormatSettings(move_footnotes=True, include_contents=True, ocr_language="nl")
    out = pdf_text(pipeline.load(paper, s).export("pdf", s))
    assert "Inhoud" in out and "Noten" in out and "reader [Noot 1]" in out
    assert "[Note 1]" not in out and "Contents" not in out
    assert "Individual differences are discussed in Section 4." in out


def test_page_map_links_converted_pages_to_their_original_pages(paper):
    """The side-by-side view turns the other side along using this map."""
    session = pipeline.load(paper)
    data = session.export("pdf", FormatSettings())
    with pymupdf.open(stream=data, filetype="pdf") as d:
        n_conv = len(d)
    with pymupdf.open(paper) as d:
        n_orig = len(d)
    m = session.page_map
    assert m and all(0 <= k < n_conv for k in m)
    assert all(0 <= p < n_orig for v in m.values() for p in v)
    assert {p for v in m.values() for p in v} == set(range(n_orig))  # every original page is reachable
    firsts = [min(m[k]) for k in sorted(m)]
    assert firsts == sorted(firsts)  # reading order: going forward never goes back in the original


def test_page_map_of_a_book_spread_uses_physical_pages(samples):
    session = pipeline.load(samples / "book_spread.pdf")
    session.export("pdf", FormatSettings())
    with pymupdf.open(samples / "book_spread.pdf") as d:
        n = len(d)
    assert all(0 <= p < n for v in session.page_map.values() for p in v)


def test_other_exports_keep_the_pdf_page_map(paper):
    session = pipeline.load(paper)
    session.export("pdf", FormatSettings())
    before = dict(session.page_map)
    session.export("docx", FormatSettings())
    session.export("printable_pdf", FormatSettings())
    assert session.page_map == before
