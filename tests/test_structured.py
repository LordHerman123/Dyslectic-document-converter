"""Word (.docx) and EPUB files as input: their own structure is read, and the original is shown as pages."""
import io
import zipfile

from PIL import Image

from dyslexia_converter import pipeline
from dyslexia_converter.model import BlockKind
from dyslexia_converter.render import preview
from dyslexia_converter.settings import FormatSettings


def _png(w=60, h=40) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (w, h), (200, 60, 60)).save(out, "PNG")
    return out.getvalue()


def _docx(path):
    from docx import Document

    d = Document()
    d.core_properties.title = "Plants"
    d.add_heading("Photosynthesis", 0)
    d.add_paragraph("Plants turn light into energy. " * 8)
    d.add_heading("What plants need", 1)
    p = d.add_paragraph("Three things: ")
    p.add_run("light").bold = True
    p.add_run(" and ")
    p.add_run("water").italic = True
    for t in ("Light", "Water"):
        d.add_paragraph(t, style="List Bullet")
    for t in ("First step", "Second step"):
        d.add_paragraph(t, style="List Number")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "In", "Out", "Water", "Oxygen"
    d.add_picture(io.BytesIO(_png()))
    d.add_heading("Why it matters", 2)
    for _ in range(30):
        d.add_paragraph("Almost all life depends on it, directly or indirectly. " * 4)
    d.save(path)


def _epub(path):
    chapter = """<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>x</title>
<style>p { color: black }</style></head><body>
<h1>Chapter One</h1>
<p>It was a <b>bright</b> cold day in <i>April</i>.<span epub:type="pagebreak" title="7">7</span> The clocks
   were striking.</p>
<ol><li>Wake up</li><li>Read</li></ol>
<blockquote><p>A quoted line.</p></blockquote>
<figure><img src="images/pic.png" alt=""/><figcaption>A red picture</figcaption></figure>
<table><tr><th>Name</th><th>Age</th></tr><tr><td>Ann</td><td>9</td></tr></table>
<h2>Part two</h2><p>The end of the chapter.</p>
</body></html>"""
    opf = """<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>A Small Book</dc:title><dc:creator>A. Writer</dc:creator>
<dc:language>en</dc:language><dc:identifier id="id">x</dc:identifier></metadata>
<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
<item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/>
<item id="pic" href="images/pic.png" media-type="image/png"/></manifest>
<spine><itemref idref="nav"/><itemref idref="c1"/></spine></package>"""
    nav = """<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><body><nav><ol><li>Contents entry</li>
</ol></nav></body></html>"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml", """<?xml version="1.0"?><container version="1.0"
 xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf"
 media-type="application/oebps-package+xml"/></rootfiles></container>""")
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/nav.xhtml", nav)
        z.writestr("OEBPS/c1.xhtml", chapter)
        z.writestr("OEBPS/images/pic.png", _png())


def test_word_file_is_read_with_its_own_structure(tmp_path):
    path = tmp_path / "plants.docx"
    _docx(path)
    s = pipeline.load(path, FormatSettings())
    doc = s.document
    kinds = [(b.kind, b.level, b.text) for b in doc.blocks]
    assert (BlockKind.TITLE, 0, "Photosynthesis") in kinds
    assert (BlockKind.HEADING, 1, "What plants need") in kinds
    assert (BlockKind.HEADING, 2, "Why it matters") in kinds
    items = [b.text for b in doc.blocks if b.kind == BlockKind.LIST_ITEM]
    assert items == ["• Light", "• Water", "1. First step", "2. Second step"]
    para = next(b for b in doc.blocks if b.text.startswith("Three things"))
    assert [(para.text[r.start:r.end], r.bold, r.italic) for r in para.styles] == [("light", True, False),
                                                                                   ("water", False, True)]
    table = next(b for b in doc.blocks if b.kind == BlockKind.TABLE)
    assert table.table.rows == [["In", "Out"], ["Water", "Oxygen"]]
    assert any(b.kind == BlockKind.IMAGE for b in doc.blocks)
    # the original is shown as pages, and the blocks know their page in it (the last ones are further on)
    assert s.original_pdf and preview.page_count(s.original_pdf) == len(doc.pages) >= 2
    assert doc.blocks[0].page == 0 and doc.blocks[-1].page > 0
    assert preview.render_page(s.original_pdf, 0, 200)[:4] == b"\x89PNG"
    for fmt in ("pdf", "docx", "epub"):
        assert s.export(fmt, FormatSettings())
    assert path.read_bytes()[:2] == b"PK"  # the source is untouched


def test_epub_book_is_read_in_reading_order(tmp_path):
    path = tmp_path / "book.epub"
    _epub(path)
    s = pipeline.load(path, FormatSettings())
    doc = s.document
    assert doc.title == "A Small Book" and doc.author == "A. Writer"
    texts = [(b.kind, b.text) for b in doc.blocks]
    assert texts[0] == (BlockKind.HEADING, "Chapter One")
    first = doc.blocks[1]
    assert first.text == "It was a bright cold day in April. The clocks were striking."  # page marker left out
    assert [first.text[r.start:r.end] for r in first.styles] == ["bright", "April"]
    assert (BlockKind.LIST_ITEM, "1. Wake up") in texts and (BlockKind.LIST_ITEM, "2. Read") in texts
    assert (BlockKind.QUOTE, "A quoted line.") in texts
    assert (BlockKind.CAPTION, "A red picture") in texts
    assert any(b.kind == BlockKind.IMAGE for b in doc.blocks)
    assert next(b for b in doc.blocks if b.kind == BlockKind.TABLE).table.rows == [["Name", "Age"], ["Ann", "9"]]
    assert (BlockKind.HEADING, "Part two") in texts
    assert not any("Contents entry" in b.text for b in doc.blocks)  # the book's own table of contents is skipped
    assert s.original_pdf and s.export("pdf", FormatSettings())


def _with_extra(src, dst, extra: dict, rename: dict | None = None):
    """A copy of an EPUB with extra files (and chapters renamed)."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for n in zin.namelist():
            data = zin.read(n)
            for old, new in (rename or {}).items():
                data = data.replace(old.encode(), new.encode())
                n = n.replace(old.replace("%20", " "), new.replace("%20", " "))
            zout.writestr(n, data)
        for n, data in extra.items():
            zout.writestr(n, data)


def test_copy_protected_books_are_named_as_such(tmp_path):
    """A shop's DRM (Adobe, Apple) cannot be read: the reader is told why, instead of a technical error."""
    import pytest

    from dyslexia_converter.extract.structured import ProtectedFile, protection

    book = tmp_path / "book.epub"
    _epub(book)
    adobe = tmp_path / "adobe.epub"
    _with_extra(book, adobe, {"META-INF/rights.xml": "<adept:rights xmlns:adept='http://ns.adobe.com/adept'/>"})
    enc = ("<encryption xmlns='urn:oasis:names:tc:opendocument:xmlns:container'><EncryptedData "
           "xmlns='http://www.w3.org/2001/04/xmlenc#'><EncryptionMethod Algorithm='{alg}'/><CipherData>"
           "<CipherReference URI='{uri}'/></CipherData></EncryptedData></encryption>")
    locked = tmp_path / "locked.epub"
    _with_extra(book, locked, {"META-INF/encryption.xml": enc.format(
        alg="http://www.w3.org/2001/04/xmlenc#aes128-cbc", uri="OEBPS/c1.xhtml")})
    fonts = tmp_path / "fonts.epub"  # only the fonts are scrambled (common in DRM-free books): readable
    _with_extra(book, fonts, {"META-INF/encryption.xml": enc.format(
        alg="http://www.idpf.org/2008/embedding", uri="OEBPS/fonts/a.otf")})
    assert protection(adobe) == "Adobe DRM" and protection(locked) == "DRM"
    assert protection(fonts) is None and protection(book) is None
    for path in (adobe, locked):
        with pytest.raises(ProtectedFile) as err:
            pipeline.load(path, FormatSettings())
        assert err.value.kind == "drm"
    assert pipeline.load(fonts, FormatSettings()).document.blocks
    word = tmp_path / "secret.docx"
    word.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600)  # a Word file saved with a password
    with pytest.raises(ProtectedFile) as err:
        pipeline.load(word, FormatSettings())
    assert err.value.kind == "password"


def test_a_book_the_layout_engine_cannot_open_still_reads(tmp_path, monkeypatch):
    """When PyMuPDF cannot lay out an (unprotected) book, the Original view shows its text in plain pages."""
    from dyslexia_converter.extract import structured

    book = tmp_path / "book.epub"
    _epub(book)

    def refuse(path):
        raise RuntimeError("Failed to open file as type epub")

    monkeypatch.setattr(structured, "original_pdf", refuse)
    s = pipeline.load(book, FormatSettings())
    assert s.original_pdf and preview.page_count(s.original_pdf) >= 1
    assert any(b.text == "Chapter One" for b in s.document.blocks)


def test_chapter_names_with_spaces_are_found(tmp_path):
    """Chapters listed as "c%201.xhtml" are the file "c 1.xhtml"."""
    book, spaced = tmp_path / "book.epub", tmp_path / "spaced.epub"
    _epub(book)
    _with_extra(book, spaced, {}, rename={"c1.xhtml": "c%201.xhtml"})
    s = pipeline.load(spaced, FormatSettings())
    assert any(b.text == "Chapter One" for b in s.document.blocks)


def test_pdf_loading_is_unchanged(paper):
    """A PDF still goes through the PDF reader, and is shown as itself."""
    s = pipeline.load(paper, FormatSettings())
    assert s.original_pdf is None and s.document.blocks
