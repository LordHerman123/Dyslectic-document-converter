"""DOCX export (python-docx).

Word supports font, size, line and paragraph spacing, letter spacing,
margins and headings (which appear in Word's Navigation pane as a document
map). Extra word spacing has no Word equivalent and is not applied. Fonts
are referenced by name, not embedded.
"""
from __future__ import annotations

import io

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..settings import FormatSettings
from .compose import ComposeResult, RItem

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def _set_char_spacing(run, pts: float) -> None:
    if not pts:
        return
    rpr = run._r.get_or_add_rPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(int(round(pts * 20))))  # twentieths of a point
    rpr.append(sp)


def _shade(paragraph, hex_fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    ppr.append(shd)


def _add_runs(p, item: RItem, s: FormatSettings, size: float, bold_all: bool = False) -> None:
    color = RGBColor.from_string(s.text_color.lstrip("#").upper()) if not s.ink_saving else RGBColor(0, 0, 0)
    for r in item.runs:
        run = p.add_run(r.text)
        run.font.name = s.font
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), s.font)
        run.font.size = Pt(size)
        run.font.bold = r.bold or bold_all
        run.font.italic = r.italic
        run.font.superscript = r.superscript and not r.marker
        run.font.color.rgb = color
        _set_char_spacing(run, s.letter_spacing)


def build_docx(result: ComposeResult, s: FormatSettings, title: str = "", author: str = "") -> bytes:
    d = DocxDocument()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(s.margin_top), Cm(s.margin_bottom)
    avail = 21.0 - s.margin_left - s.margin_right
    col = min(avail, s.reading_width)
    extra = (avail - col) / 2
    sec.left_margin = Cm(s.margin_left + extra)
    sec.right_margin = Cm(s.margin_right + extra)
    d.core_properties.title = title or "Converted document"
    d.core_properties.author = author
    d.core_properties.comments = "Reformatted for easier reading by Dyslexia Converter"

    normal = d.styles["Normal"]
    normal.font.name = s.font
    normal.font.size = Pt(s.font_size)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = s.line_spacing
    pf.space_after = Pt(s.paragraph_spacing)
    pf.alignment = ALIGN.get(s.alignment, WD_ALIGN_PARAGRAPH.LEFT)
    for lvl in range(1, 5):
        st = d.styles[f"Heading {lvl}"]
        st.font.name = s.font
        st.font.size = Pt(s.font_size * {1: 1.15, 2: 1.05}.get(lvl, 1.0) * s.heading_scale)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(s.paragraph_spacing * 1.4)
        st.paragraph_format.space_after = Pt(s.paragraph_spacing * 0.6)
    tst = d.styles["Title"]
    tst.font.name = s.font
    tst.font.size = Pt(s.font_size * 1.45 * s.heading_scale)
    tst.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

    items = result.items
    i = 0
    while i < len(items):
        it = items[i]
        k = it.kind
        if k == "title":
            p = d.add_paragraph(style="Title")
            _add_runs(p, it, s, s.font_size * 1.45 * s.heading_scale, bold_all=True)
        elif k in ("heading", "box_heading"):
            lvl = max(1, min(4, it.level or 1))
            p = d.add_heading(level=lvl)
            _add_runs(p, it, s, s.font_size * {1: 1.15, 2: 1.05}.get(lvl, 1.0) * s.heading_scale, bold_all=True)
        elif k == "image" and it.image is not None:
            w = min(Cm(col).emu, int((it.natural_width or 300) * 1.35 * 12700))
            d.add_picture(io.BytesIO(it.image.data), width=w)
            d.paragraphs[-1].paragraph_format.keep_with_next = True
        elif k == "table":
            tab = it.table
            if tab is not None and tab.reliable and s.table_mode != "image":
                n = max(len(r) for r in tab.rows)
                t = d.add_table(rows=0, cols=n)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.LEFT
                for ri, row in enumerate(tab.rows):
                    cells = t.add_row().cells
                    for ci in range(n):
                        cp = cells[ci].paragraphs[0]
                        run = cp.add_run(row[ci] if ci < len(row) else "")
                        run.font.size = Pt(max(9.0, s.font_size * 0.8))
                        run.font.name = s.font
                        run.bold = ri == 0
                        cp.paragraph_format.line_spacing = 1.2
                        cp.paragraph_format.space_after = Pt(2)
                d.add_paragraph()
            elif tab is not None and tab.fallback_image is not None:
                w = min(Cm(col).emu, int((it.natural_width or 300) * 1.35 * 12700))
                d.add_picture(io.BytesIO(tab.fallback_image.data), width=w)
                note = d.add_paragraph().add_run("Table shown as an image of the original.")
                note.font.size = Pt(max(8.5, s.font_size * 0.78))
                note.font.name = s.font
        else:
            p = d.add_paragraph()
            size = s.font_size
            if k in ("caption", "small", "about", "footnote"):
                size = max(8.5, s.font_size * 0.82)
            elif k in ("reference", "endnote"):
                size = s.font_size * 0.93
            if it.marker:
                p.paragraph_format.left_indent = Cm(1.0)
                p.paragraph_format.first_line_indent = Cm(-1.0)
                mr = p.add_run(it.marker + "\t")
                mr.font.size = Pt(size)
                mr.font.name = s.font
            _add_runs(p, it, s, size)
            if k in ("box_paragraph", "quote") and s.boxed_sections and not s.ink_saving:
                _shade(p, "F3F0E2")
            if k == "caption" and it.keep_with_next:
                p.paragraph_format.keep_with_next = True
        i += 1

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
